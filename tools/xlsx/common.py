"""Shared helpers for the xlsx tools: LibreOffice (headless, local) and formula utilities.

Everything here runs locally. Nothing is uploaded anywhere.
"""

import os
import re
import shutil
import subprocess
import tempfile
import warnings
from pathlib import Path

warnings.filterwarnings("ignore", module="openpyxl")

import openpyxl  # noqa: E402
from openpyxl.formula.tokenizer import Token, Tokenizer  # noqa: E402
from openpyxl.utils import column_index_from_string, get_column_letter  # noqa: E402
from openpyxl.worksheet.formula import ArrayFormula  # noqa: E402

RECALC_MACRO = """<?xml version="1.0" encoding="UTF-8"?>
<!DOCTYPE script:module PUBLIC "-//OpenOffice.org//DTD OfficeDocument 1.0//EN" "module.dtd">
<script:module xmlns:script="http://openoffice.org/2000/script" script:name="Module1" script:language="StarBasic">
    Sub RecalculateAndSave()
      ThisComponent.calculateAll()
      ThisComponent.store()
      ThisComponent.close(True)
    End Sub
</script:module>"""


class LibreOfficeError(RuntimeError):
    pass


# --------------------------------------------------------------------------- LibreOffice

def find_soffice():
    path = shutil.which("soffice") or shutil.which("libreoffice")
    if not path:
        raise LibreOfficeError("LibreOffice (soffice) is not installed. " + install_hint())
    return path


def install_hint():
    return ("Install it with: apt-get install -y --no-install-recommends libreoffice-calc "
            "(the 'core' package alone cannot open spreadsheets).")


def diagnose():
    """Return a list of human-readable findings about the local LibreOffice install."""
    notes = []
    exe = shutil.which("soffice")
    if not exe:
        return ["soffice not found. " + install_hint()]
    try:
        out = subprocess.run([exe, "--version"], capture_output=True, text=True, timeout=60)
        notes.append("version: " + (out.stdout.strip() or out.stderr.strip()))
    except subprocess.TimeoutExpired:
        notes.append("soffice --version timed out")
    prog = Path(exe).resolve().parent
    if (prog / "libsclo.so").exists() or (prog / "libsclo.dylib").exists():
        notes.append("Calc (spreadsheet) component: present")
    elif prog.exists() and any(prog.glob("lib*.so")):
        notes.append("Calc (spreadsheet) component: MISSING. " + install_hint())
    return notes


def _env():
    env = os.environ.copy()
    env["SAL_USE_VCLPLUGIN"] = "svp"
    return env


def _new_profile(tmpdir, timeout):
    profile = Path(tmpdir) / "lo_profile"
    url = profile.as_uri()
    try:
        subprocess.run([find_soffice(), "--headless", "--terminate_after_init",
                        f"-env:UserInstallation={url}"],
                       capture_output=True, timeout=timeout, env=_env())
    except subprocess.TimeoutExpired:
        raise LibreOfficeError("LibreOffice timed out creating a profile. Diagnosis: "
                               + "; ".join(diagnose()))
    macro_dir = profile / "user" / "basic" / "Standard"
    if not macro_dir.exists():
        raise LibreOfficeError("LibreOffice did not create a usable profile. Diagnosis: "
                               + "; ".join(diagnose()))
    (macro_dir / "Module1.xba").write_text(RECALC_MACRO)
    return url


def recalc_in_place(path, timeout=600):
    """Recalculate every formula in an .xlsx with LibreOffice and save it in place.

    Use it on a temporary copy: LibreOffice rewrites the whole file.
    """
    path = Path(path).resolve()
    before = (path.stat().st_mtime_ns, path.stat().st_size)
    with tempfile.TemporaryDirectory(prefix="xlsx_tools_") as tmp:
        url = _new_profile(tmp, min(timeout, 120))
        cmd = [find_soffice(), "--headless", "--norestore", f"-env:UserInstallation={url}",
               "vnd.sun.star.script:Standard.Module1.RecalculateAndSave?language=Basic&location=application",
               str(path)]
        try:
            res = subprocess.run(cmd, capture_output=True, text=True, timeout=timeout, env=_env())
        except subprocess.TimeoutExpired:
            raise LibreOfficeError(f"Recalculation timed out after {timeout}s. Diagnosis: "
                                   + "; ".join(diagnose()))
    if res.returncode != 0:
        raise LibreOfficeError(f"LibreOffice exited {res.returncode}: {res.stderr.strip()}")
    if (path.stat().st_mtime_ns, path.stat().st_size) == before:
        raise LibreOfficeError("LibreOffice ran but did not rewrite the file (nothing recalculated). "
                               "Diagnosis: " + "; ".join(diagnose()))


def convert(path, target, outdir, timeout=600):
    """Convert a file with LibreOffice (e.g. target='pdf:calc_pdf_Export:{...}'). Returns output path."""
    path = Path(path).resolve()
    outdir = Path(outdir).resolve()
    outdir.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="xlsx_tools_") as tmp:
        url = _new_profile(tmp, min(timeout, 120))
        cmd = [find_soffice(), "--headless", "--norestore", f"-env:UserInstallation={url}",
               "--convert-to", target, "--outdir", str(outdir), str(path)]
        try:
            res = subprocess.run(cmd, capture_output=True, text=True, timeout=timeout, env=_env())
        except subprocess.TimeoutExpired:
            raise LibreOfficeError(f"Conversion timed out after {timeout}s. Diagnosis: "
                                   + "; ".join(diagnose()))
    ext = target.split(":")[0]
    out = outdir / (path.stem + "." + ext)
    if res.returncode != 0 or not out.exists():
        raise LibreOfficeError(f"Conversion failed ({res.returncode}): {res.stderr.strip() or res.stdout.strip()}")
    return out


# --------------------------------------------------------------------------- workbook helpers

def load(path, values=False):
    """Load a workbook. values=True gives cached values (formulas gone) - never save that one."""
    return openpyxl.load_workbook(path, data_only=values)


def formula_text(value):
    """Return the formula string ("=...") held in a cell value, or None."""
    if isinstance(value, ArrayFormula):
        return value.text if str(value.text).startswith("=") else "=" + str(value.text)
    if isinstance(value, str) and value.startswith("="):
        return value
    return None


def is_array_formula(value):
    return isinstance(value, ArrayFormula)


def iter_cells(ws):
    """Existing cells in row-major order (skips the empty grid openpyxl would otherwise create)."""
    for key in sorted(ws._cells):
        yield ws._cells[key]


# --------------------------------------------------------------------------- formula patterns

_REF_PART = re.compile(r"^(\$?)([A-Za-z]{1,3})?(\$?)(\d+)?$")


def _split_sheet(ref):
    if "!" not in ref:
        return "", ref
    if ref.startswith("'"):
        end = ref.find("'!")
        while end != -1 and ref[end - 1:end + 1] == "''":
            end = ref.find("'!", end + 1)
        if end != -1:
            return ref[:end + 2], ref[end + 2:]
    idx = ref.rfind("!")
    return ref[:idx + 1], ref[idx + 1:]


def _part_to_r1c1(part, row, col):
    m = _REF_PART.match(part)
    if not m or (m.group(2) is None and m.group(4) is None):
        return None
    cabs, cletters, rabs, rdigits = m.groups()
    out = ""
    if rdigits is not None:
        r = int(rdigits)
        out += f"R{r}" if rabs else f"R[{r - row}]"
    if cletters is not None:
        c = column_index_from_string(cletters.upper())
        out += f"C{c}" if cabs else f"C[{c - col}]"
    return out


def to_r1c1(formula, row, col):
    """Rewrite an A1 formula as relative R1C1 so copies of the same formula look identical.

    `=B3*2` in C3 and `=B4*2` in C4 both become `=R[0]C[-1]*2`.
    """
    try:
        tok = Tokenizer(formula)
    except Exception:
        return formula
    for t in tok.items:
        if t.type == Token.OPERAND and t.subtype == Token.RANGE:
            sheet, ref = _split_sheet(t.value)
            parts = ref.split(":")
            if len(parts) == 1 and not re.match(r"^\$?[A-Za-z]{1,3}\$?\d+$", parts[0]):
                continue  # a defined name such as PE or S_MIN, not a cell
            conv = [_part_to_r1c1(p, row, col) for p in parts]
            if all(c is not None for c in conv):
                t.value = sheet + ":".join(conv)
    try:
        return "=" + tok.render().lstrip("=")
    except Exception:
        return formula


def compress_ranges(coords):
    """Turn a list of cell coordinates into a short readable summary like 'C4:C76, E2'."""
    from openpyxl.utils.cell import coordinate_from_string
    cells = []
    for c in coords:
        letters, r = coordinate_from_string(c)
        cells.append((column_index_from_string(letters), r))
    cells = sorted(set(cells))
    # group by column into vertical runs, then merge identical runs across adjacent columns
    runs = {}
    for c, r in cells:
        runs.setdefault(c, []).append(r)
    col_runs = []
    for c in sorted(runs):
        rows = runs[c]
        start = prev = rows[0]
        for r in rows[1:]:
            if r == prev + 1:
                prev = r
                continue
            col_runs.append((c, start, prev))
            start = prev = r
        col_runs.append((c, start, prev))
    blocks = []
    for c, r1, r2 in col_runs:
        if blocks and blocks[-1][1] == c - 1 and blocks[-1][2] == r1 and blocks[-1][3] == r2:
            blocks[-1][1] = c
        else:
            blocks.append([c, c, r1, r2])
    parts = []
    for c1, c2, r1, r2 in blocks:
        a = f"{get_column_letter(c1)}{r1}"
        b = f"{get_column_letter(c2)}{r2}"
        parts.append(a if a == b else f"{a}:{b}")
    return ", ".join(parts)
