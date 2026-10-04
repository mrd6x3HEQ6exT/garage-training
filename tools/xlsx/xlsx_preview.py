#!/usr/bin/env python3
"""Picture preview: render each sheet of a workbook to a PNG image (LibreOffice, local only).

The picture is close to what Excel shows - fills, borders, fonts, column widths,
conditional formatting - but it is LibreOffice's drawing, not Excel's. Dropdown arrows,
cell checkboxes and Excel-only features will not look identical.

Usage:
  python3 xlsx_preview.py FILE.xlsx [-o OUTDIR] [--sheets "Name 1" "Name 2"]
         [--range A1:Z40] [--max-px 2400] [--include-hidden] [--no-recalc]

--range     only draw this block of the chosen sheet(s) (works by hiding everything else
            in a throwaway copy, so the original file is never touched)
--max-px    longest side of each image in pixels
--no-recalc skip recalculating formulas first (faster; only safe when the file was last
            saved by Excel)
Output: OUTDIR/NN_<sheet>.png (default ./xlsx_out/<file name>_preview/).
"""

import argparse
import re
import shutil
import subprocess
import sys
import tempfile
from copy import copy
from pathlib import Path

from openpyxl.utils import get_column_letter
from openpyxl.worksheet.dimensions import ColumnDimension
from openpyxl.utils.cell import range_boundaries

from common import LibreOfficeError, convert, load, recalc_in_place

PDF_ONE_PAGE_PER_SHEET = 'pdf:calc_pdf_Export:{"SinglePageSheets":{"type":"boolean","value":"true"}}'


def _split_column_blocks(ws):
    """A width saved for a block of columns (e.g. C:CH) would override hiding single columns,
    so give every column in a block its own entry first."""
    for key, d in list(ws.column_dimensions.items()):
        if d.min and d.max and d.max > d.min:
            for c in range(d.min, d.max + 1):
                letter = get_column_letter(c)
                nd = ColumnDimension(ws, index=letter, width=d.width, bestFit=d.bestFit, hidden=d.hidden,
                                     outlineLevel=d.outlineLevel, collapsed=d.collapsed,
                                     customWidth=d.customWidth)
                nd._style = copy(d._style)
                nd.min = nd.max = c
                ws.column_dimensions[letter] = nd


def _hide_outside(path, sheets, cell_range):
    min_col, min_row, max_col, max_row = range_boundaries(cell_range)
    wb = load(path)
    for name in sheets:
        ws = wb[name]
        _split_column_blocks(ws)
        last_row = max(ws.max_row, max_row)
        last_col = max(ws.max_column, max_col)
        for r in list(range(1, min_row)) + list(range(max_row + 1, last_row + 1)):
            ws.row_dimensions[r].hidden = True
        for c in list(range(1, min_col)) + list(range(max_col + 1, last_col + 1)):
            ws.column_dimensions[get_column_letter(c)].hidden = True
    wb.save(path)


def _page_count(pdf):
    out = subprocess.run(["pdfinfo", str(pdf)], capture_output=True, text=True).stdout
    m = re.search(r"Pages:\s+(\d+)", out)
    return int(m.group(1)) if m else 0


def preview(file, outdir=None, sheets=None, cell_range=None, max_px=2400, include_hidden=False, recalc=True):
    if not shutil.which("pdftoppm"):
        raise SystemExit("pdftoppm is missing (apt-get install -y poppler-utils)")
    file = Path(file)
    outdir = Path(outdir) if outdir else Path("xlsx_out") / f"{file.stem}_preview"
    outdir.mkdir(parents=True, exist_ok=True)
    wb = load(file)
    names = wb.sheetnames
    states = {ws.title: ws.sheet_state for ws in wb.worksheets}
    if sheets:
        missing = [s for s in sheets if s not in names]
        if missing:
            raise SystemExit(f"No such sheet(s): {missing}. Sheets are: {names}")
    if cell_range and not sheets:
        raise SystemExit("--range needs --sheets (which sheet(s) to crop)")
    wanted = sheets or [n for n in names if include_hidden or states[n] == "visible"]

    with tempfile.TemporaryDirectory(prefix="xlsx_preview_") as tmp:
        work = Path(tmp) / "preview.xlsx"
        shutil.copyfile(file, work)
        work.chmod(0o644)
        if cell_range:
            _hide_outside(work, sheets, cell_range)
        if recalc:
            recalc_in_place(work)
        pdf = convert(work, PDF_ONE_PAGE_PER_SHEET, tmp)
        pages = _page_count(pdf)
        written = []
        if pages != len(names):
            print(f"warning: {pages} pages for {len(names)} sheets - names may not line up; "
                  "writing page_NN.png files instead", file=sys.stderr)
            for p in range(1, pages + 1):
                target = outdir / f"page_{p:02d}"
                subprocess.run(["pdftoppm", "-png", "-singlefile", "-scale-to", str(max_px),
                                "-f", str(p), "-l", str(p), str(pdf), str(target)], check=True)
                written.append(target.with_suffix(".png"))
            return written
        for i, name in enumerate(names, 1):
            if name not in wanted:
                continue
            safe = re.sub(r"[^A-Za-z0-9._-]+", "_", name)
            target = outdir / f"{i:02d}_{safe}"
            subprocess.run(["pdftoppm", "-png", "-singlefile", "-scale-to", str(max_px),
                            "-f", str(i), "-l", str(i), str(pdf), str(target)], check=True)
            written.append(target.with_suffix(".png"))
    return written


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("file")
    ap.add_argument("-o", "--outdir")
    ap.add_argument("--sheets", nargs="+")
    ap.add_argument("--range", dest="cell_range", help="e.g. A1:Z40 (needs --sheets)")
    ap.add_argument("--max-px", type=int, default=2400)
    ap.add_argument("--include-hidden", action="store_true")
    ap.add_argument("--no-recalc", action="store_true")
    args = ap.parse_args(argv)
    try:
        files = preview(args.file, args.outdir, args.sheets, args.cell_range, args.max_px,
                        args.include_hidden, not args.no_recalc)
    except LibreOfficeError as e:
        print(f"LibreOffice problem: {e}", file=sys.stderr)
        return 2
    for f in files:
        print(f)
    return 0


if __name__ == "__main__":
    sys.exit(main())
