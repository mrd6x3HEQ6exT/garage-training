#!/usr/bin/env python3
"""Rules test: prove a gradebook's GO/NG and percentage formulas follow its own rules.

The rules are read from the workbook itself - never typed into this script:
  * layout "rows" : a hidden "Rules" sheet (Tab, Section, Step #, Type, Max, Critical, Part,
                    Minimum group, Switch) plus its minimum-group and GO-floor tables
  * layout "v9"   : the original layout (students in columns, steps in rows). Rules are parsed
                    from the overall GO / NG formula; the red step numbers and any
                    "Booth GO?" / "Report GO?" rows are read as extra views to compare against.

`run` builds test students from those rules (perfect, blank, one step missing, every
critical step missed, every S step under the minimum, every section exactly at and one
under its minimum, invalid values, the lowest passing score, switches off, random mixes),
types them into a throwaway copy, recalculates with LibreOffice, and checks every result
cell against what the rules say. The original file is never modified.

Usage:
  python3 rules_test.py spec FILE.xlsx [--layout auto|v9|rows] [-o spec.json]
  python3 rules_test.py diff A.xlsx|A.json B.xlsx|B.json
  python3 rules_test.py run  FILE.xlsx [--layout ...] [--pes "TAB" ...] [-o report.md] [--keep DIR]

Exit code of `run` is 1 when any check fails. Runs locally only; specs and reports contain
course data, so keep them out of git (the folder's .gitignore covers xlsx_out/ and *.json).
"""

import argparse
import json
import random
import re
import shutil
import sys
import tempfile
from pathlib import Path

from openpyxl.utils import column_index_from_string, get_column_letter

from common import LibreOfficeError, formula_text, load, recalc_in_place

PCT_TOL = 1e-6


# =========================================================================== spec model

def _step(sid, section, num, typ, mx, critical=False, part=None, group=None, switch=None, loc=None):
    return {"id": sid, "section": section, "num": num, "type": typ, "max": mx, "critical": critical,
            "part": part, "group": group, "switch": switch, "loc": loc}


def _label(step):
    return f"{step['section']} step {step['num']}"


def _norm(s):
    return re.sub(r"\s+", " ", str(s or "").replace("\n", " ")).strip()


# --------------------------------------------------------------------------- rows layout

def _find_blocks(header):
    """Locate the three tables in the Rules sheet by their header names."""
    idx = {}
    for i, h in enumerate(header):
        idx.setdefault(_norm(h).lower(), []).append(i)
    def first_after(name, start):
        for i in idx.get(name, []):
            if i >= start:
                return i
        return None
    steps = {k: first_after(k, 0) for k in ("tab", "section", "step #", "type", "max", "critical", "part",
                                            "minimum group", "switch", "column")}
    g_tab = first_after("tab", (steps["switch"] or 0) + 1)
    groups = {"tab": g_tab, "group": first_after("minimum group", g_tab or 0),
              "minimum": first_after("minimum", g_tab or 0)}
    f_tab = first_after("tab", (groups["minimum"] or 0) + 1)
    floor = {"tab": f_tab, "floor": first_after("go floor", f_tab or 0), "max": first_after("max pts", f_tab or 0)}
    return steps, groups, floor


def spec_from_rules_sheet(wb, path):
    ws = wb["Rules"]
    rows = list(ws.iter_rows(values_only=True))
    header = rows[0]
    st, gr, fl = _find_blocks(header)
    s_min = 8
    for r in rows[:3]:
        for i, v in enumerate(r):
            if _norm(v).lower().startswith("scalable (s) step minimum") and i + 1 < len(r) and isinstance(r[i + 1], (int, float)):
                s_min = int(r[i + 1])
    spec = {"source": str(path), "layout": "rows", "s_min": s_min, "pes": {}}
    for r in rows[1:]:
        tab = r[st["tab"]]
        if not tab or r[st["step #"]] is None:
            continue
        pe = spec["pes"].setdefault(tab, {"steps": [], "groups": {}, "floor": 0, "stated_total": None,
                                          "stated_part_max": {}, "views": {}, "notes": []})
        sid = len(pe["steps"])
        typ = _norm(r[st["type"]]).upper()
        pe["steps"].append(_step(sid, _norm(r[st["section"]]), r[st["step #"]], typ, r[st["max"]],
                                 _norm(r[st["critical"]]).upper() in ("Y", "YES", "1", "TRUE", "★"),
                                 _norm(r[st["part"]]) or None, _norm(r[st["minimum group"]]) or None,
                                 _norm(r[st["switch"]]) or None, _norm(r[st["column"]])))
    for r in rows[1:]:
        tab = r[gr["tab"]] if gr["tab"] is not None else None
        if tab and tab in spec["pes"] and r[gr["group"]]:
            pe = spec["pes"][tab]
            name = _norm(r[gr["group"]])
            members = [s["id"] for s in pe["steps"] if s["group"] == name]
            pe["groups"][name] = {"steps": members, "min": int(r[gr["minimum"]])}
        tab = r[fl["tab"]] if fl["tab"] is not None else None
        if tab and tab in spec["pes"]:
            pe = spec["pes"][tab]
            pe["floor"] = float(r[fl["floor"]] or 0) if fl["floor"] is not None else 0
            pe["stated_total"] = r[fl["max"]] if fl["max"] is not None else None
    for tab, pe in spec["pes"].items():
        true_total = sum(s["max"] for s in pe["steps"])
        if pe["stated_total"] not in (None, true_total):
            pe["notes"].append(f"stated max points {pe['stated_total']} but steps add up to {true_total}")
        orphan = {s["group"] for s in pe["steps"] if s["group"]} - set(pe["groups"])
        if orphan:
            pe["notes"].append(f"steps name minimum group(s) with no minimum set: {sorted(orphan)}")
    return spec


# --------------------------------------------------------------------------- v9 layout

_SUM_DIV = re.compile(r"SUM\(\s*C(\d+)\s*:\s*C(\d+)\s*\)\s*/\s*(\d+(?:\.\d+)?)")


def _split_top(s):
    out, depth, cur, q = [], 0, "", False
    for ch in s:
        if ch == '"':
            q = not q
        if not q:
            if ch == "(":
                depth += 1
            elif ch == ")":
                depth -= 1
            elif ch == "," and depth == 0:
                out.append(cur.strip())
                cur = ""
                continue
        cur += ch
    if cur.strip():
        out.append(cur.strip())
    return out


def _and_args(formula):
    """Arguments of the AND(...) that leads to "GO"."""
    f = formula.replace("\n", " ")
    i = f.find("IF(AND(")
    if i < 0:
        return None
    i += len("IF(AND(")
    depth, j = 1, i
    while j < len(f) and depth:
        if f[j] == "(":
            depth += 1
        elif f[j] == ")":
            depth -= 1
        j += 1
    return _split_top(f[i:j - 1])


def _rows_of(refs):
    rows = []
    for part in refs.split(","):
        part = part.strip().replace("$", "")
        m = re.fullmatch(r"C(\d+)(?::C(\d+))?", part)
        if not m:
            return None
        a = int(m.group(1))
        b = int(m.group(2) or a)
        rows += list(range(a, b + 1))
    return rows


def _parse_clause(ws, row, row_to_step, notes, seen=None):
    """Parse the GO condition in column C of `row` into critical steps, groups and S minimums."""
    seen = seen or set()
    if row in seen:
        return {"critical": set(), "groups": [], "smin": {}, "unparsed": [], "refs": set()}
    seen.add(row)
    f = formula_text(ws.cell(row, 3).value) or ""
    res = {"critical": set(), "groups": [], "smin": {}, "unparsed": [], "refs": set()}
    args = _and_args(f)
    if args is None:
        res["unparsed"].append(f"row {row}: no IF(AND(...),\"GO\") found")
        return res
    for a in args:
        t = re.sub(r"\s+", "", a)
        m = re.fullmatch(r'C(\d+)="GO"', t)
        if m:
            sub = _parse_clause(ws, int(m.group(1)), row_to_step, notes, seen)
            for k in ("critical", "refs"):
                res[k] |= sub[k]
            res["groups"] += sub["groups"]
            res["smin"].update(sub["smin"])
            res["unparsed"] += sub["unparsed"]
            continue
        m = re.fullmatch(r"SUM\(([^)]*)\)(>=|>|=)(\d+)", t)
        if m:
            rows = _rows_of(m.group(1))
            n = int(m.group(3))
            if rows is None or any(r not in row_to_step for r in rows):
                res["unparsed"].append(a)
                continue
            ids = [row_to_step[r] for r in rows]
            res["refs"] |= set(ids)
            if m.group(2) == "=":
                if n == len(ids):
                    res["critical"] |= set(ids)
                else:
                    res["unparsed"].append(a)
            else:
                res["groups"].append({"steps": ids, "min": n + (1 if m.group(2) == ">" else 0)})
            continue
        m = re.fullmatch(r"(C\d+(?::C\d+)?)(>=|=)(\d+)", t)
        if m:
            rows = _rows_of(m.group(1))
            n = int(m.group(3))
            if rows is None or any(r not in row_to_step for r in rows):
                res["unparsed"].append(a)
                continue
            ids = [row_to_step[r] for r in rows]
            res["refs"] |= set(ids)
            if n == 1:
                res["critical"] |= set(ids)
            else:
                for i in ids:
                    res["smin"][i] = n
            continue
        res["unparsed"].append(a)
    return res


def _dv_types(ws):
    types = {}
    for dv in ws.data_validations.dataValidation:
        f1 = str(dv.formula1 or "").replace(" ", "")
        if dv.type == "list" and f1.strip('"') == "0,8,9,10":
            t = "S"
        elif dv.type == "whole" and str(dv.formula1) == "0" and str(dv.formula2) == "1":
            t = "NS"
        else:
            continue
        for rng in dv.sqref.ranges:
            if rng.min_col <= 3 <= rng.max_col:
                for r in range(rng.min_row, rng.max_row + 1):
                    types[r] = t
    return types


def _is_red(cell):
    c = cell.font.color if cell.font else None
    return c is not None and c.type == "rgb" and str(c.rgb).upper().endswith("FF0000")


def v9_pe_tabs(wb):
    return [ws for ws in wb.worksheets if _norm(ws.cell(2, 1).value).upper().startswith("GO / NG")]


def _v9_step_rows(ws):
    rows = {}
    r = 4
    while r <= ws.max_row:
        if isinstance(ws.cell(r, 2).value, (int, float)):
            rows[r] = len(rows)
        elif rows:
            break
        r += 1
    return rows


def _dashboard_rows(wb, tab):
    """Row numbers of this tab that other sheets read (directly or through INDIRECT)."""
    used = set()
    pat = re.compile(r"\$?C\$?(\d+)(?::\$?[A-Z]{1,3}\$?\d+)?")
    for ws in wb.worksheets:
        if ws.title == tab or _norm(ws.cell(2, 1).value).upper().startswith("GO / NG"):
            continue
        for row in ws.iter_rows():
            for c in row:
                f = formula_text(c.value)
                if not f:
                    continue
                if f"'{tab}'!" in f or f"{tab}!" in f:
                    for m in re.finditer(re.escape(f"'{tab}'!") + r"\$?C\$?(\d+)", f):
                        used.add(int(m.group(1)))
                if "INDIRECT(" in f:
                    # INDIRECT("'"&$E6&"'!$C$99:$CH$99"): the tab name sits in E6 on this sheet
                    for ind in re.finditer(r'INDIRECT\((.*?)\)', f):
                        body = ind.group(1)
                        names = [ws[f"{m.group(1)}{m.group(2)}"].value
                                 for m in re.finditer(r"&\$?([A-Z]{1,3})\$?(\d+)&", body)]
                        if tab in names:
                            used |= {int(x) for x in pat.findall(body)}
    return used


def v9_outputs(ws, row_to_step=None, dashboard=None):
    """Every result row on a v9 tab: GO/NG, Final %, and each pass/fail or percentage row
    below the steps (labelled or not), with the steps it covers."""
    row_to_step = row_to_step if row_to_step is not None else _v9_step_rows(ws)
    dashboard = dashboard or set()
    last = max(row_to_step) if row_to_step else 3
    outs = [{"kind": "result", "row": 2, "label": "GO / NG (row 2)"},
            {"kind": "final", "row": 3, "label": "Final % (row 3)"}]
    for r in range(last + 1, ws.max_row + 1):
        f = formula_text(ws.cell(r, 3).value)
        if not f:
            continue
        lab = _norm(ws.cell(r, 1).value) or _norm(ws.cell(r, 2).value)
        lab = lab if re.search(r"[A-Za-z%]", lab) else "unlabeled"
        label = f"{lab} (row {r}" + (", read by a dashboard)" if r in dashboard else ")")
        if "IF(AND(" in f.replace(" ", "") and '"GO"' in f:
            c = _parse_clause(ws, r, row_to_step, [])
            if c["refs"]:
                outs.append({"kind": "scope_result", "row": r, "label": label,
                             "scope": sorted(c["refs"]), "critical": sorted(c["critical"])})
            continue
        m = _SUM_DIV.search(f)
        if m:
            a, b, n = int(m.group(1)), int(m.group(2)), float(m.group(3))
            scope = sorted(row_to_step[x] for x in range(a, b + 1) if x in row_to_step)
            outs.append({"kind": "scope_pct", "row": r, "label": label, "scope": scope,
                         "divisor": n, "range": (a, b)})
    return outs


def spec_from_v9(wb, path):
    spec = {"source": str(path), "layout": "v9", "s_min": 8, "pes": {}}
    for ws in v9_pe_tabs(wb):
        pe = {"steps": [], "groups": {}, "floor": 0, "stated_total": None, "stated_part_max": {},
              "views": {}, "notes": [], "rows": {}}
        dv = _dv_types(ws)
        section = ""
        row_to_step = {}
        r = 4
        while r <= ws.max_row:
            num = ws.cell(r, 2).value
            if not isinstance(num, (int, float)):
                if pe["steps"]:
                    break
                r += 1
                continue
            if ws.cell(r, 1).value:
                section = _norm(ws.cell(r, 1).value).replace(" (", " (")
            label_type = "S" if re.search(r"\(S\)", section) else ("NS" if "(NS)" in section else None)
            typ = dv.get(r) or label_type or "NS"
            sid = len(pe["steps"])
            pe["steps"].append(_step(sid, section, int(num), typ, 10 if typ == "S" else 1, loc=r))
            row_to_step[r] = sid
            r += 1
        no_dv = [s["loc"] for s in pe["steps"] if s["loc"] not in dv]
        if no_dv:
            from common import compress_ranges
            pe["notes"].append("no dropdown/validation on step rows " +
                               compress_ranges([f"C{x}" for x in no_dv]).replace("C", "") + " (any number can be typed)")
        mism = [s["loc"] for s in pe["steps"] if s["loc"] in dv and
                ((re.search(r"\(S\)", s["section"]) and dv[s["loc"]] != "S") or ("(NS)" in s["section"] and dv[s["loc"]] != "NS"))]
        if mism:
            pe["notes"].append(f"dropdown type disagrees with the section label on rows {mism}")

        clause = _parse_clause(ws, 2, row_to_step, pe["notes"])
        for sid in clause["critical"]:
            pe["steps"][sid]["critical"] = True
        for g in clause["groups"]:
            first = pe["steps"][g["steps"][0]]["section"]
            last = pe["steps"][g["steps"][-1]]["section"]
            name = first if first == last else f"{first} to {last}"
            if name in pe["groups"]:
                name = f"{name} #{len(pe['groups']) + 1}"
            pe["groups"][name] = g
            for sid in g["steps"]:
                pe["steps"][sid]["group"] = name
        if clause["unparsed"]:
            pe["notes"].append("overall formula terms not understood: " + "; ".join(clause["unparsed"]))
        s_ids = {s["id"] for s in pe["steps"] if s["type"] == "S"}
        if set(clause["smin"]) != s_ids:
            miss = s_ids - set(clause["smin"])
            extra = set(clause["smin"]) - s_ids
            if miss:
                pe["notes"].append("scalable steps with no 8+ check in the overall formula: "
                                   + ", ".join(_label(pe["steps"][i]) for i in sorted(miss)))
            if extra:
                pe["notes"].append("8+ check on non-scalable steps: "
                                   + ", ".join(_label(pe["steps"][i]) for i in sorted(extra)))
        mins = set(clause["smin"].values())
        if mins and mins != {8}:
            pe["notes"].append(f"scalable minimum is not 8 everywhere: {sorted(mins)}")

        # Final % and the part percentage rows tell us the maxima and the booth/report split
        f3 = formula_text(ws.cell(3, 3).value) or ""
        m = _SUM_DIV.search(f3)
        if m:
            pe["stated_total"] = float(m.group(3))
        fm = re.search(r"MAX\(\s*(0?\.\d+|1(?:\.0+)?)\s*,", f3)
        if fm:
            pe["floor"] = float(fm.group(1))
        outs = v9_outputs(ws, row_to_step, _dashboard_rows(wb, ws.title))
        pe["outputs"] = outs
        all_ids = [s["id"] for s in pe["steps"]]
        for o in outs:
            if o["kind"] != "scope_pct":
                continue
            a, b = o["range"]
            n = o["divisor"]
            pe["stated_part_max"][o["label"]] = n
            if o["scope"] and len(o["scope"]) < len(all_ids):
                lab = o["label"].lower()
                if "booth" in lab or lab.startswith("b%"):
                    part = "Booth"
                elif "report" in lab or lab.startswith("r%"):
                    part = "Report"
                else:
                    part = "Booth" if 0 in o["scope"] else "Report"
                for sid in o["scope"]:
                    if pe["steps"][sid]["part"] is None:
                        pe["steps"][sid]["part"] = part
            real = sum(pe["steps"][i]["max"] for i in o["scope"])
            if n != real:
                pe["notes"].append(f"{o['label']} divides by {n:g} but those steps are worth {real}")
            beyond = [x for x in range(a, b + 1) if x not in row_to_step]
            if beyond:
                pe["notes"].append(f"{o['label']} adds rows {a}-{b}, which includes {len(beyond)} row(s) that are not steps")
        true_total = sum(s["max"] for s in pe["steps"])
        if pe["stated_total"] is not None and pe["stated_total"] != true_total:
            pe["notes"].append(f"Final % divides by {pe['stated_total']:g} but the steps are worth {true_total}")

        # extra views to compare with the overall rules: red step numbers and every
        # pass/fail row below the steps (Booth GO?, Report GO?, unlabeled helpers)
        pe["views"]["red step numbers"] = {
            "critical": sorted(s["id"] for s in pe["steps"] if _is_red(ws.cell(s["loc"], 2))), "scope": all_ids}
        for o in outs:
            if o["kind"] == "scope_result":
                pe["views"][o["label"]] = {"critical": o["critical"], "scope": o["scope"]}
        spec["pes"][ws.title] = pe
    return spec


def load_spec(path, layout="auto"):
    p = Path(path)
    if p.suffix.lower() == ".json":
        return json.loads(p.read_text())
    wb = load(p)
    if layout == "auto":
        layout = "rows" if "Rules" in wb.sheetnames else ("v9" if v9_pe_tabs(wb) else None)
    if layout == "rows":
        return spec_from_rules_sheet(wb, p)
    if layout == "v9":
        return spec_from_v9(wb, p)
    raise SystemExit("Could not tell the layout; pass --layout v9 or --layout rows")


# =========================================================================== expected results

def allowed(step, v):
    return v in ((0, 8, 9, 10) if step["type"] == "S" else (0, 1))


def expected(pe, s_min, values, switches, layout, out):
    """What the rules say one output cell should show for one test student.
    Returns None when this case does not pin the cell down (it is then not checked)."""
    steps = [s for s in pe["steps"] if not (s["switch"] and switches.get(s["switch"]) is False)]
    included = {s["id"] for s in steps}
    vals = {s["id"]: values.get(s["id"]) for s in steps}
    entered = [s for s in steps if vals[s["id"]] is not None]
    invalid = any(not allowed(s, vals[s["id"]]) for s in entered)

    def judge(sub):
        if any(s["type"] == "S" and vals[s["id"]] < s_min for s in sub):
            return "NG"
        if any(s["critical"] and vals[s["id"]] < 1 for s in sub):
            return "NG"
        ids = {s["id"] for s in sub}
        for g in pe["groups"].values():
            live = [i for i in g["steps"] if i in included]
            if not live or not all(i in ids for i in live):
                continue  # group switched off, or not part of this slice
            if sum(vals[i] for i in live) < g["min"]:
                return "NG"
        return "GO"

    kind = out["kind"]
    if kind in ("result", "final"):
        if not entered:
            return ""
        if invalid:
            return "CHECK" if (layout == "rows" and kind == "result") else ("" if layout == "rows" else None)
        if len(entered) < len(steps):
            return ("INC" if layout == "rows" else "") if kind == "result" else ""
        res = judge(steps)
        if kind == "result":
            return res
        pct = sum(vals[s["id"]] for s in steps) / sum(s["max"] for s in steps)
        return max(pe["floor"], pct) if (res == "GO" and pe["floor"]) else pct

    # a slice of the steps: booth / report parts, or the steps a v9 helper row covers
    if "scope" in out:
        sub = [pe["steps"][i] for i in out["scope"] if i in included]
    else:
        sub = [s for s in steps if s["part"] == out["part"]]
    if layout == "rows":
        if not entered:
            return None
        if invalid:
            return ""
        if not sub:
            return "n/a"
        if any(vals[s["id"]] is None for s in sub):
            return ""
    else:
        # v9 helper rows often wait for the WHOLE tab before showing anything (their blank
        # check covers every step) - that is a design choice, so only judge complete students
        if not sub or len(entered) < len(steps) or invalid:
            return None
    if kind in ("part_result", "scope_result"):
        return judge(sub)
    return sum(vals[s["id"]] for s in sub) / sum(s["max"] for s in sub)


# =========================================================================== test cases

def make_cases(pe, s_min, layout, seed=7):
    steps = pe["steps"]
    switch_names = sorted({s["switch"] for s in steps if s["switch"]})
    on = {n: True for n in switch_names}

    def perfect(sw=on):
        return {s["id"]: (10 if s["type"] == "S" else 1) for s in steps
                if not (s["switch"] and sw.get(s["switch"]) is False)}

    cases = [("perfect score", perfect(), on), ("nothing entered", {}, on)]
    v = perfect()
    v.pop(steps[-1]["id"], None)
    cases.append((f"{_label(steps[-1])} left blank", v, on))
    for s in steps:
        if s["critical"]:
            v = perfect()
            v[s["id"]] = 0
            cases.append((f"critical {_label(s)} = 0", v, on))
    for s in steps:
        if s["type"] == "S":
            v = perfect()
            v[s["id"]] = 0
            cases.append((f"scalable {_label(s)} = 0 (under {s_min})", v, on))
    v = perfect()
    for s in steps:
        if s["type"] == "S":
            v[s["id"]] = s_min
    cases.append((f"every scalable step at {s_min}", v, on))
    lowest = dict(v)
    for name, g in pe["groups"].items():
        noncrit = [i for i in g["steps"] if not steps[i]["critical"]]
        k = len(g["steps"]) - g["min"]
        if k >= 1 and len(noncrit) >= k:
            v = perfect()
            for i in noncrit[:k]:
                v[i] = 0
            cases.append((f"{name}: exactly at minimum ({g['min']} of {len(g['steps'])})", v, on))
            for i in noncrit[:k]:
                lowest[i] = 0
        if len(noncrit) >= k + 1:
            v = perfect()
            for i in noncrit[:k + 1]:
                v[i] = 0
            cases.append((f"{name}: one under minimum", v, on))
    cases.append(("lowest passing score", lowest, on))
    if layout == "rows":
        s_steps = [s for s in steps if s["type"] == "S"]
        ns_steps = [s for s in steps if s["type"] == "NS"]
        if s_steps:
            v = perfect()
            v[s_steps[0]["id"]] = 5
            cases.append((f"invalid 5 in {_label(s_steps[0])}", v, on))
        if ns_steps:
            v = perfect()
            v[ns_steps[0]["id"]] = 2
            cases.append((f"invalid 2 in {_label(ns_steps[0])}", v, on))
    rnd = random.Random(seed)
    for k in range(3):
        v = {}
        for s in steps:
            if s["type"] == "S":
                v[s["id"]] = rnd.choice([10, 10, 9, 9, 8, 0])
            else:
                v[s["id"]] = 1 if rnd.random() < 0.85 else 0
        cases.append((f"random mix #{k + 1}", v, on))
    for name in switch_names:
        sw = dict(on)
        sw[name] = False
        cases.append((f"switch {name} off, the rest perfect", perfect(sw), sw))
        v = perfect(sw)
        crit = [s for s in steps if s["critical"] and s["switch"] != name]
        if crit:
            v[crit[0]["id"]] = 0
            cases.append((f"switch {name} off, critical {_label(crit[0])} = 0", v, sw))
    return cases


# =========================================================================== layout adapters

class V9Layout:
    name = "v9"
    first_col = 3

    def __init__(self, wb):
        self.wb = wb

    def capacity(self, tab):
        ws = self.wb[tab]
        last = 3
        for dv in ws.data_validations.dataValidation:
            for rng in dv.sqref.ranges:
                last = max(last, rng.max_col)
        return last - self.first_col + 1

    def outputs(self, tab, pe):
        return pe.get("outputs") or v9_outputs(self.wb[tab])

    def write(self, wb, tab, slot, pe, values, switches):
        ws = wb[tab]
        col = self.first_col + slot
        for s in pe["steps"]:
            ws.cell(s["loc"], col).value = values.get(s["id"])

    def read(self, wbv, tab, slot, out):
        return wbv[tab].cell(out["row"], self.first_col + slot).value


class RowsLayout:
    name = "rows"
    first_row = 10
    header_row = 9

    def __init__(self, wb):
        self.wb = wb
        self.roster = wb["Roster"] if "Roster" in wb.sheetnames else None

    def capacity(self, tab):
        ws = self.wb[tab]
        n = 0
        r = self.first_row
        while formula_text(ws.cell(r, 1).value):
            n += 1
            r += 1
        return n

    def _header_cols(self, tab):
        ws = self.wb[tab]
        cols = {}
        for c in range(1, ws.max_column + 1):
            h = _norm(ws.cell(self.header_row, c).value).upper().replace(" | ", " ")
            cols.setdefault(h, c)
        return cols

    def outputs(self, tab, pe):
        cols = self._header_cols(tab)
        want = [("result", None, "RESULT"), ("final", None, "FINAL %"),
                ("part_pct", "Booth", "BOOTH %"), ("part_pct", "Report", "REPORT %"),
                ("part_result", "Booth", "BOOTH GO/NG"), ("part_result", "Report", "REPORT GO/NG")]
        outs = []
        for kind, part, h in want:
            if h in cols:
                outs.append({"kind": kind, "part": part, "col": cols[h],
                             "label": f"{h.title()} (column {get_column_letter(cols[h])})"})
        return outs

    def switch_cells(self, tab):
        ws = self.wb[tab]
        cells = {}
        for r in range(1, self.header_row):
            lab = re.sub(r"\s+", "", _norm(ws.cell(r, 1).value)).upper()
            if lab:
                cells[lab] = ws.cell(r, 2).coordinate
        return cells

    def write(self, wb, tab, slot, pe, values, switches):
        ws = wb[tab]
        row = self.first_row + slot
        if self.roster is not None:
            wb["Roster"].cell(2 + slot, 1).value = f"Test student {slot + 1:03d}"
        for s in pe["steps"]:
            ws[f"{s['loc']}{row}"] = values.get(s["id"])
        cells = self.switch_cells(tab)
        for name, state in switches.items():
            key = re.sub(r"\s+", "", name).upper()
            for lab, coord in cells.items():
                if lab.startswith(key):
                    ws[coord] = "Enabled" if state else "Disabled"

    def read(self, wbv, tab, slot, out):
        return wbv[tab].cell(self.first_row + slot, out["col"]).value


# =========================================================================== run

def _same(exp, act):
    if exp is None:
        return True  # not checked for this case
    if isinstance(exp, float):
        return isinstance(act, (int, float)) and abs(act - exp) < PCT_TOL
    act = "" if act is None else act
    return str(exp) == str(act)


def _show(v):
    if v is None:
        return "(blank)"
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return f"{v * 100:.2f}%"
    return repr(v) if v != "" else "(blank)"


def run(file, layout="auto", pes=None, keep=None, max_fail=12):
    spec = load_spec(file, layout)
    layout = spec["layout"]
    wb0 = load(file)
    adapter = V9Layout(wb0) if layout == "v9" else RowsLayout(wb0)
    tabs = [t for t in spec["pes"] if (not pes or t in pes) and t in wb0.sheetnames]
    plan = {}
    for tab in tabs:
        pe = spec["pes"][tab]
        cases = make_cases(pe, spec["s_min"], layout)
        cap = adapter.capacity(tab)
        groups = {}
        for c in cases:
            groups.setdefault(json.dumps(c[2], sort_keys=True), []).append(c)
        chunks = []
        for g in groups.values():
            for i in range(0, len(g), cap):
                chunks.append(g[i:i + cap])
        plan[tab] = chunks
    n_batches = max((len(v) for v in plan.values()), default=0)
    results = {tab: [] for tab in tabs}
    tmpdir = Path(keep) if keep else Path(tempfile.mkdtemp(prefix="rules_test_"))
    tmpdir.mkdir(parents=True, exist_ok=True)
    try:
        for b in range(n_batches):
            work = tmpdir / f"batch_{b + 1:02d}.xlsx"
            wb = load(file)
            for tab, chunks in plan.items():
                if b >= len(chunks):
                    continue
                for slot, (name, values, sw) in enumerate(chunks[b]):
                    adapter.write(wb, tab, slot, spec["pes"][tab], values, sw)
            wb.save(work)
            recalc_in_place(work)
            wbv = load(work, values=True)
            for tab, chunks in plan.items():
                if b >= len(chunks):
                    continue
                pe = spec["pes"][tab]
                outs = adapter.outputs(tab, pe)
                for slot, (name, values, sw) in enumerate(chunks[b]):
                    for out in outs:
                        e = expected(pe, spec["s_min"], values, sw, layout, out)
                        act = adapter.read(wbv, tab, slot, out)
                        results[tab].append((name, out["label"], e, act, _same(e, act)))
    finally:
        if not keep:
            shutil.rmtree(tmpdir, ignore_errors=True)
    return spec, results


def report_md(spec, results, max_fail=12):
    lines = [f"# Rules test - {Path(spec['source']).name}", "",
             f"Layout: **{spec['layout']}**. Rules read from: "
             + ("the hidden Rules sheet" if spec["layout"] == "rows" else
                "each tab's overall GO / NG formula (row 2)") + f". Scalable minimum: {spec['s_min']}.", "",
             "| Tab | Test students | Checks | Failed |", "|---|---|---|---|"]
    total_fail = 0
    for tab, rows in results.items():
        names = {r[0] for r in rows}
        checked = [r for r in rows if r[2] is not None]
        failed = [r for r in checked if not r[4]]
        total_fail += len(failed)
        lines.append(f"| {tab} | {len(names)} | {len(checked)} | {len(failed) or '-'} |")
    lines += ["", f"**{total_fail} failed check(s).**" if total_fail else "**All checks passed.**"]
    for tab, rows in results.items():
        pe = spec["pes"][tab]
        failed = [r for r in rows if r[2] is not None and not r[4]]
        notes = pe.get("notes", [])
        views = _view_diffs(pe)
        if not (failed or notes or views):
            continue
        lines += ["", f"## {tab}"]
        for n in notes:
            lines.append(f"- Note: {n}")
        for v in views:
            lines.append(f"- {v}")
        by_out = {}
        for name, label, e, a, ok in failed:
            by_out.setdefault(label, []).append((name, e, a))
        for label, items in by_out.items():
            lines.append(f"- **{label}** wrong for {len(items)} test student(s):")
            for name, e, a in items[:max_fail]:
                lines.append(f"  - {name}: rules say {_show(e)}, sheet shows {_show(a)}")
            if len(items) > max_fail:
                lines.append(f"  - ... {len(items) - max_fail} more")
    return "\n".join(lines) + "\n", total_fail


def _view_diffs(pe):
    """Compare critical-step lists from other places in the file with the rules used."""
    out = []
    crit = {s["id"] for s in pe["steps"] if s["critical"]}
    for view, v in pe.get("views", {}).items():
        ids = set(v["critical"])
        ref = crit & set(v["scope"])
        missing = sorted(ref - ids)
        extra = sorted(ids - ref)
        if missing:
            out.append(f"{view}: does NOT treat as critical: "
                       + ", ".join(_label(pe["steps"][i]) for i in missing))
        if extra:
            out.append(f"{view}: treats as critical but the overall rules do not: "
                       + ", ".join(_label(pe["steps"][i]) for i in extra))
    return out


# =========================================================================== spec / diff output

def spec_summary(spec):
    lines = [f"# Rules found in {Path(spec['source']).name} (layout {spec['layout']}, scalable minimum {spec['s_min']})", ""]
    for tab, pe in spec["pes"].items():
        steps = pe["steps"]
        crit = [s for s in steps if s["critical"]]
        lines.append(f"## {tab}: {len(steps)} steps, {sum(s['max'] for s in steps)} points, "
                     f"{len(crit)} critical" + (f", GO floor {pe['floor']:.0%}" if pe["floor"] else ""))
        bysec = {}
        for s in crit:
            bysec.setdefault(s["section"], []).append(str(s["num"]))
        if bysec:
            lines.append("- critical: " + "; ".join(f"{k} {', '.join(v)}" for k, v in bysec.items()))
        for name, g in pe["groups"].items():
            lines.append(f"- {name}: need {g['min']} of {len(g['steps'])}")
        sw = sorted({s["switch"] for s in steps if s["switch"]})
        if sw:
            lines.append(f"- switches: {', '.join(sw)}")
        for n in pe.get("notes", []):
            lines.append(f"- Note: {n}")
        for v in _view_diffs(pe):
            lines.append(f"- {v}")
        lines.append("")
    return "\n".join(lines)


def diff_specs(a, b):
    lines = [f"# Rules: {Path(a['source']).name}  vs  {Path(b['source']).name}", ""]
    same = True
    for tab in sorted(set(a["pes"]) | set(b["pes"])):
        pa, pb = a["pes"].get(tab), b["pes"].get(tab)
        if pa is None or pb is None:
            lines.append(f"- {tab}: only in {'second' if pa is None else 'first'} file")
            same = False
            continue
        out = []
        if len(pa["steps"]) != len(pb["steps"]):
            out.append(f"step count {len(pa['steps'])} vs {len(pb['steps'])}")
        part_diff, renamed = [], {}
        for sa, sb in zip(pa["steps"], pb["steps"]):
            if sa["num"] != sb["num"]:
                out.append(f"step order differs at {_label(sa)} vs {_label(sb)}")
                break
            if sa["section"] != sb["section"]:
                renamed.setdefault((sa["section"], sb["section"]), 0)
                renamed[(sa["section"], sb["section"])] += 1
            for k, lab in (("type", "type"), ("max", "max")):
                if sa[k] != sb[k]:
                    out.append(f"{_label(sa)}: {lab} {sa[k]} vs {sb[k]}")
            if sa["critical"] != sb["critical"]:
                out.append(f"{_label(sa)}: critical in {'first' if sa['critical'] else 'second'} file only")
            if sa["part"] and sb["part"] and sa["part"] != sb["part"]:
                part_diff.append(sa)
        for (x, y), n in renamed.items():
            out.append(f"section name '{x}' vs '{y}' ({n} steps; names only, rules compared anyway)")
        if part_diff:
            out.append(f"booth/report split differs on {len(part_diff)} step(s), e.g. "
                       + ", ".join(f"{_label(s)}" for s in part_diff[:5]))
        ga = sorted((tuple(g["steps"]), g["min"]) for g in pa["groups"].values())
        gb = sorted((tuple(g["steps"]), g["min"]) for g in pb["groups"].values())
        if ga != gb:
            for g in set(ga) ^ set(gb):
                where = "first" if g in ga else "second"
                out.append(f"section minimum only in {where}: need {g[1]} of "
                           + f"{pa['steps'][g[0][0]]['section'] if g[0][0] < len(pa['steps']) else '?'} ({len(g[0])} steps)")
        if (pa["floor"] or 0) != (pb["floor"] or 0):
            out.append(f"GO floor {pa['floor']} vs {pb['floor']}")
        swa = sorted({s["switch"] for s in pa["steps"] if s["switch"]})
        swb = sorted({s["switch"] for s in pb["steps"] if s["switch"]})
        if swa != swb:
            out.append(f"switches {swa} vs {swb}")
        if out:
            same = False
            lines.append(f"## {tab}")
            lines += [f"- {o}" for o in out]
    if same:
        lines.append("Rules are identical.")
    return "\n".join(lines) + "\n"


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    p = sub.add_parser("spec", help="show the rules found in a workbook")
    p.add_argument("file")
    p.add_argument("--layout", default="auto", choices=["auto", "v9", "rows"])
    p.add_argument("-o", "--output", help="also save the rules as JSON (keep it out of git)")
    p = sub.add_parser("diff", help="compare the rules of two workbooks (or saved JSON specs)")
    p.add_argument("a")
    p.add_argument("b")
    p = sub.add_parser("run", help="test the formulas against the rules")
    p.add_argument("file")
    p.add_argument("--layout", default="auto", choices=["auto", "v9", "rows"])
    p.add_argument("--pes", nargs="+", help="only these tabs")
    p.add_argument("-o", "--output", help="write the report to this .md file")
    p.add_argument("--keep", help="keep the filled-in test workbooks in this folder")
    p.add_argument("--max-failures", type=int, default=12, help="examples shown per wrong cell")
    args = ap.parse_args(argv)

    if args.cmd == "spec":
        spec = load_spec(args.file, args.layout)
        print(spec_summary(spec))
        if args.output:
            Path(args.output).write_text(json.dumps(spec, indent=1))
            print(f"Saved {args.output}")
        return 0
    if args.cmd == "diff":
        print(diff_specs(load_spec(args.a), load_spec(args.b)))
        return 0
    try:
        spec, results = run(args.file, args.layout, args.pes, args.keep, args.max_failures)
    except LibreOfficeError as e:
        print(f"LibreOffice problem: {e}", file=sys.stderr)
        return 2
    text, fails = report_md(spec, results, args.max_failures)
    if args.output:
        Path(args.output).write_text(text)
        print(f"Wrote {args.output}: {'all checks passed' if not fails else f'{fails} failed check(s)'}")
    else:
        print(text)
    return 1 if fails else 0


if __name__ == "__main__":
    sys.exit(main())
