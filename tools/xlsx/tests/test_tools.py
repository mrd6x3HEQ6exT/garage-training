"""Self-tests for the xlsx tools. They build small made-up workbooks - no real data.

Run from tools/xlsx:   python3 -m unittest discover -s tests -v
Tests that need LibreOffice / pdftoppm are skipped when those are not installed.
"""

import shutil
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent))

import openpyxl  # noqa: E402
from openpyxl.styles import Font  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402

import rules_test  # noqa: E402
import xlsx_compare  # noqa: E402
import xlsx_inspect  # noqa: E402
import xlsx_security  # noqa: E402
from common import to_r1c1  # noqa: E402

HAVE_LO = shutil.which("soffice") is not None
HAVE_PDFTOPPM = shutil.which("pdftoppm") is not None


def make_demo(path, booth_divisor=13, booth_row_checks_critical=False):
    """A tiny gradebook in the original layout: students in columns C-E, steps in rows 4-8.

    Alpha (S) step 1 (row 4) must be 8+; Alpha (NS) steps 1-4 (rows 5-8) need 3 of 4 and
    step 4 is critical (red). Row 10 'Booth %' and row 11 'Booth GO?' carry deliberate bugs
    unless told otherwise: divisor 13 instead of 14, and step 4 not treated as critical.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Demo PE 01"
    ws["A1"] = "Demo PE 01"
    ws["A2"] = "GO / NG:"
    ws["A3"] = "Final %:"
    ws["A4"] = "Alpha\n(S)"
    ws["A5"] = "Alpha\n(NS)"
    for r, n in zip(range(4, 9), [1, 1, 2, 3, 4]):
        ws.cell(r, 2, n)
    ws["B8"].font = Font(color="FFFF0000")
    ws["A10"] = "Booth %"
    ws["A11"] = "Booth GO?"
    for col in "CDE":
        crit = f", {col}8=1" if booth_row_checks_critical else ""
        ws[f"{col}1"] = f"Student {col}"
        ws[f"{col}2"] = (f'=IF(COUNTBLANK({col}4:{col}8)>0,"",IF(AND({col}4>=8, SUM({col}5:{col}8)>=3, '
                         f'{col}8=1),"GO","NG"))')
        ws[f"{col}3"] = f'=IF({col}2="","",SUM({col}4:{col}8)/14)'
        ws[f"{col}10"] = f'=IF(COUNTBLANK({col}4:{col}8)>0,"",SUM({col}4:{col}8)/{booth_divisor})'
        ws[f"{col}11"] = (f'=IF(COUNTBLANK({col}4:{col}8)>0,"",IF(AND({col}4>=8, SUM({col}5:{col}8)>=3{crit}),'
                          f'"GO","NG"))')
    s = DataValidation(type="list", formula1='"0,8,9,10"')
    s.add("C4:E4")
    ns = DataValidation(type="whole", operator="between", formula1="0", formula2="1")
    ns.add("C5:E8")
    ws.add_data_validation(s)
    ws.add_data_validation(ns)
    ws.protection.sheet = True
    wb.save(path)
    return path


class Base(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp(prefix="xlsx_tools_test_"))

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)


class TestCommon(unittest.TestCase):
    def test_r1c1_groups_copies(self):
        self.assertEqual(to_r1c1("=B3*2", 3, 3), to_r1c1("=B4*2", 4, 3))
        self.assertNotEqual(to_r1c1("=B3*2", 3, 3), to_r1c1("=B3*2", 4, 3))

    def test_r1c1_keeps_names(self):
        self.assertIn("S_MIN", to_r1c1('=COUNTIF(A1:A9,"<"&S_MIN)', 1, 2))
        self.assertIn("PE", to_r1c1("=SUM(PE)", 1, 1))


class TestSecurity(Base):
    def test_clean_file_passes(self):
        f = make_demo(self.tmp / "demo.xlsx")
        rep = xlsx_security.check_file(str(f))
        self.assertFalse(rep.failed, rep.items)

    def test_bad_parts_fail(self):
        f = make_demo(self.tmp / "demo.xlsx")
        bad = self.tmp / "bad.xlsx"
        with zipfile.ZipFile(f) as zin, zipfile.ZipFile(bad, "w") as zout:
            for item in zin.infolist():
                data = zin.read(item.filename)
                if item.filename == "xl/worksheets/sheet1.xml":
                    data = data.replace(b"<f>", b"<f>WEBSERVICE(&quot;http://x&quot;)+", 1)
                zout.writestr(item, data)
            zout.writestr("xl/vbaProject.bin", b"fake")
            zout.writestr("xl/externalLinks/externalLink1.xml", b"<externalLink/>")
        rep = xlsx_security.check_file(str(bad))
        failed = {i["check"] for i in rep.items if i["status"] == "FAIL"}
        self.assertIn("Macros (VBA)", failed)
        self.assertIn("Links to other workbooks", failed)
        self.assertIn("Function WEBSERVICE", failed)

    def test_password_fails(self):
        f = self.tmp / "pw.xlsx"
        make_demo(f)
        wb = openpyxl.load_workbook(f)
        wb.active.protection.password = "secret"
        wb.save(f)
        rep = xlsx_security.check_file(str(f))
        self.assertIn("Protection passwords", {i["check"] for i in rep.items if i["status"] == "FAIL"})
        self.assertFalse(xlsx_security.check_file(str(f), allow_passwords=True).failed)


class TestInspectCompare(Base):
    def test_inspect_writes_patterns(self):
        f = make_demo(self.tmp / "demo.xlsx")
        out = xlsx_inspect.inspect(f, self.tmp / "insp")
        text = (out / "01_Demo_PE_01.txt").read_text()
        self.assertIn("FORMULA PATTERNS", text)
        self.assertIn("x3", text)  # each row formula copied to C, D, E
        self.assertIn("C2:E2", text)

    def test_compare_finds_edit(self):
        a = make_demo(self.tmp / "a.xlsx")
        b = make_demo(self.tmp / "b.xlsx", booth_divisor=14)
        report, diff = xlsx_compare.compare(a, b)
        self.assertTrue(diff)
        self.assertIn("Cells changed (3)", report)
        self.assertIn("/14", report)
        same, diff2 = xlsx_compare.compare(a, a)
        self.assertFalse(diff2)


class TestRulesSpec(Base):
    def test_spec_and_views(self):
        f = make_demo(self.tmp / "demo.xlsx")
        spec = rules_test.load_spec(f)
        pe = spec["pes"]["Demo PE 01"]
        crit = [rules_test._label(s) for s in pe["steps"] if s["critical"]]
        self.assertEqual(crit, ["Alpha (NS) step 4"])
        self.assertEqual(list(pe["groups"].values())[0]["min"], 3)
        notes = " ".join(pe["notes"])
        self.assertIn("divides by 13 but those steps are worth 14", notes)
        views = " ".join(rules_test._view_diffs(pe))
        self.assertIn("Booth GO? (row 11): does NOT treat as critical: Alpha (NS) step 4", views)
        self.assertNotIn("red step numbers", views)  # red mark matches the formula


@unittest.skipUnless(HAVE_LO, "LibreOffice not installed")
class TestRulesRun(Base):
    def test_run_finds_planted_bugs_only(self):
        f = make_demo(self.tmp / "demo.xlsx")
        spec, results = rules_test.run(f)
        rows = results["Demo PE 01"]
        failed = {(r[1].split(" (")[0]) for r in rows if r[2] is not None and not r[4]}
        self.assertEqual(failed, {"Booth %", "Booth GO?"})
        checked_main = [r for r in rows if r[1].startswith(("GO / NG", "Final %")) and r[2] is not None]
        self.assertTrue(checked_main and all(r[4] for r in checked_main))

    def test_fixed_file_passes(self):
        f = make_demo(self.tmp / "ok.xlsx", booth_divisor=14, booth_row_checks_critical=True)
        spec, results = rules_test.run(f)
        _, fails = rules_test.report_md(spec, results)
        self.assertEqual(fails, 0)


@unittest.skipUnless(HAVE_LO and HAVE_PDFTOPPM, "LibreOffice or pdftoppm not installed")
class TestPreview(Base):
    def test_png_written(self):
        import xlsx_preview
        f = make_demo(self.tmp / "demo.xlsx")
        files = xlsx_preview.preview(f, self.tmp / "pv", max_px=600)
        self.assertEqual(len(files), 1)
        self.assertTrue(files[0].exists() and files[0].stat().st_size > 1000)
        files = xlsx_preview.preview(f, self.tmp / "pv2", sheets=["Demo PE 01"], cell_range="A1:C5", max_px=600)
        self.assertTrue(files[0].exists())


if __name__ == "__main__":
    unittest.main()
